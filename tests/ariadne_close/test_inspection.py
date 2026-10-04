import io
from pathlib import Path
import pytest
from openpyxl import Workbook
from app.services.ariadne_close_intake import parse_file


def complex_book():
    book = Workbook()
    sheet = book.active
    sheet.title = "Faturamento"
    sheet["A1"] = "Relatório de energia"
    sheet.merge_cells("A1:C1")
    sheet.append([])
    sheet.append(["Descrição", "Valor faturado", "Quantidade kWh"])
    sheet.append(["Energia", "26.000,00", "100000"])
    sheet["F2"] = "=SUM(B4:B4)"
    extra = book.create_sheet("Referências")
    extra.append(["Nota", "Origem"])
    extra.append(["Contexto", "Operador"])
    hidden = book.create_sheet("Oculta")
    hidden.sheet_state = "hidden"
    hidden.append(["Valor", "Unidade"])
    hidden.append(["250", "BRL/MWh"])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def test_inspection_keeps_titles_formulas_hidden_and_regions():
    preview = parse_file(complex_book(), "normal.xlsx", inspection=True)
    assert preview["parserVersion"] == "ariadne.inspect.v2"
    table = next(t for t in preview["tables"] if t["sheet"] == "Faturamento")
    assert table["headerRow"] == 3
    assert table["rows"][0]["values"]["Valor faturado"] == "26.000,00"
    assert table["rows"][0]["cells"]["Valor faturado"] == "Faturamento!B4"
    assert any(c.get("formula") for s in preview["sheets"] for c in s["cells"])
    assert next(s for s in preview["sheets"] if s["name"] == "Oculta")["hidden"]
    assert next(t for t in preview["tables"] if t["sheet"] == "Oculta")["rows"][0][
        "issues"
    ]


def test_strict_mode_remains_strict():
    with pytest.raises(ValueError):
        parse_file(complex_book(), "normal.xlsx")


@pytest.mark.parametrize("filename", ["bad.pdf", "bad.xlsx"])
def test_known_signature_mismatch_is_specific(filename):
    from app.services.ariadne_close_intake import parse_isolated

    with pytest.raises(ValueError, match="Assinatura do conteúdo"):
        parse_isolated(b"not the declared format", filename, inspection=True)


def test_nonblank_multicell_title_does_not_hide_real_header():
    book = Workbook()
    sheet = book.active
    sheet.append(["Fechamento", "Setembro"])
    sheet.append(["item_key", "amount"])
    sheet.append(["Energia", 26])
    data = io.BytesIO()
    book.save(data)
    preview = parse_file(data.getvalue(), "title.xlsx", inspection=True)
    assert preview["tables"][0]["headerRow"] == 2
    assert preview["tables"][0]["proposedMapping"] == {
        "item_key": "item_key",
        "amount": "amount",
    }


def test_nonregion_data_remains_visible_in_coverage_inventory():
    book = Workbook()
    sheet = book.active
    sheet.append(["item_key", "amount"])
    sheet.append(["Energia", 5])
    sheet.append([])
    sheet.append(["Other billed fee", None, 10])
    data = io.BytesIO()
    book.save(data)
    preview = parse_file(data.getvalue(), "extra.xlsx", inspection=True)
    assert "Sheet!C4" in preview["sheets"][0]["unclassifiedCells"]
    assert any("completude" in w for w in preview["warnings"])


def test_uncached_formula_outside_region_prevents_claiming_complete_source():
    book = Workbook()
    sheet = book.active
    sheet.append(["item_key", "amount"])
    sheet.append(["energy", 5])
    sheet["C5"] = "=10"
    data = io.BytesIO()
    book.save(data)
    preview = parse_file(data.getvalue(), "formula.xlsx", inspection=True)
    assert "Sheet!C5" in preview["sheets"][0]["unclassifiedCells"]


def test_preheader_numeric_text_is_not_harmless_presentation():
    book = Workbook()
    sheet = book.active
    sheet.append(["Other billed fee", "10.00"])
    sheet.append([])
    sheet.append(["item_key", "amount"])
    sheet.append(["energy", 5])
    data = io.BytesIO()
    book.save(data)
    preview = parse_file(data.getvalue(), "fees.xlsx", inspection=True)
    assert "Sheet!B1" in preview["sheets"][0]["unclassifiedCells"]


def test_grouped_decimal_requires_explicit_v2_mode_preserves_raw():
    from types import SimpleNamespace
    from app.services.ariadne_close import candidates

    csv = b'item_key,component,scope,period,currency,tax_basis,amount\nenergy,energy,unit,2021-09,BRL,inclusive,"26.000,00"\n'
    source = SimpleNamespace(
        preview=parse_file(csv, "invoice.csv", inspection=True), role="invoice"
    )
    review = SimpleNamespace(scope="unit", period="2021-09")
    mapping = source.preview["tables"][0]["proposedMapping"]
    body = {"mapping": mapping, "sheet": "CSV", "numericMode": "comma"}
    row = candidates(source, review, body)[0]
    assert row["amount"] == "26000.00" and not row["errors"]
    assert row["raw"]["amount"] == "26.000,00"
    assert "grouped decimal" in row["valueOrigins"]["amount"]
    assert candidates(source, review, {**body, "numericMode": "strict"})[0]["errors"]


def test_portable_native_invoice_grammar_and_ambiguity():
    from app.services.ariadne_close_inspection import propose

    native = "\n".join(
        [
            "SYNTHETIC ENGINEERING TEXT",
            "Companhia Energética do Ceará CNPJ 07047251000170",
            "123456789",
            "A123.4567.8901.2345.6789",
            " B3 OUTROS",
            "Serviço Público TESTE",
            "(A) Contrato de Energia",
            "Custo de disponibilidade 10,00",
            "Adicional Band. Vermelha 100,000 20,000,20000",
            "Benefício Tarifário Bruto 1,00",
            "Subtotal(A) 31,00",
            "(B) Outros Encargos",
            "Benefício Tarifário Líquido 2,00-",
            "Subtotal(B) 2,00-",
            "20/10/2021 29,00",
            "1234567",
            "09/2021",
            "DESCRIÇAO TOTAL MEDIDO VALORES(R$)TARIFA (R$)",
        ]
    )

    def inspect(text):
        return propose(
            {
                "kind": "pdf",
                "tables": [],
                "pages": [{"number": 1, "text": text}],
                "extraction": "native",
            }
        )

    preview = inspect(native)
    assert preview["observations"][-1]["amount"] == "29.00"
    assert preview["observations"][1]["price"] == "0.20000"
    assert preview["observations"][1]["price_unit"] is None
    assert "TESTE" in preview["proposal"]["class"]
    assert not inspect(native + "\n" + native).get("observations")
    assert not inspect(native + "\n21/10/2021 30,00\n7654321").get("observations")
    multiple = propose(
        {
            "kind": "pdf",
            "tables": [],
            "pages": [
                {"number": 1, "text": native},
                {
                    "number": 2,
                    "text": native.replace("1234567\n09/2021", "7654321\n10/2021"),
                },
            ],
            "extraction": "native",
        }
    )
    assert not multiple.get("observations")
    assert "scope" not in multiple["proposal"] and "period" not in multiple["proposal"]
    assert multiple["extraction"] == "ambiguous / manual-only"
    unsupported = inspect(
        native.replace(
            "Custo de disponibilidade 10,00", "Cobrança não implementada 10,00"
        )
    )
    assert unsupported["observations"][-1]["operation"] is None


@pytest.mark.parametrize("inspection", [False, True])
def test_active_part_type_rejected_even_with_passive_filename(inspection):
    import zipfile

    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(complex_book())) as original, zipfile.ZipFile(
        output, "w"
    ) as changed:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/xl/passive.bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>',
                )
            changed.writestr(entry, data)
        changed.writestr("xl/passive.bin", b"not executable in test")
    with pytest.raises(ValueError, match="ativos"):
        parse_file(output.getvalue(), "renamed.xlsx", inspection=inspection)


PUBLIC = Path(r"C:\dev\ariadne-real-public-evaluation-20261003")


@pytest.mark.skipif(
    not PUBLIC.exists(), reason="authentic public corpus acquired separately"
)
def test_authentic_public_pdf_observations():
    preview = parse_file(
        (PUBLIC / "aris-invoice-131111946-physical-page-190.pdf").read_bytes(),
        "invoice.pdf",
        inspection=True,
    )
    assert preview["proposal"]["period"] == "2021-09"
    assert preview["proposal"]["scope"] == "9010675"
    rows = preview["observations"]
    assert [r["amount"] for r in rows] == [
        "76.73",
        "13.35",
        "4.89",
        "94.97",
        "-3.54",
        "-3.54",
        "91.43",
    ]
    assert rows[1]["quantity"] == "100.000"
    assert rows[1]["price"] == "0.13350"
    assert all(r["tax_basis"] == "unknown" for r in rows)


@pytest.mark.skipif(
    not PUBLIC.exists(), reason="authentic public corpus acquired separately"
)
def test_same_page_duplicate_invoice_is_ambiguous():
    from app.services.ariadne_close_inspection import propose

    text = (PUBLIC / "aris-physical-page-190-native-text.txt").read_text(
        encoding="utf-8"
    )
    preview = propose(
        {
            "kind": "pdf",
            "tables": [],
            "pages": [{"number": 1, "text": text + "\n" + text}],
            "extraction": "native",
        }
    )
    assert not preview.get("observations")
    assert any("ambíguo" in w for w in preview["warnings"])
